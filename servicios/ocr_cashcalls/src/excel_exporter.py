from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill

from src.extractor import CASH_CALL_CONCEPTS


TOTAL_MISMATCH_FILL = PatternFill(fill_type="solid", fgColor="F4CCCC")


def excel_to_bytes(data) -> bytes:
    """Genera en memoria un Excel con exactamente la tabla visible."""
    df_base = data.copy() if isinstance(data, pd.DataFrame) else pd.DataFrame(data)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_base.to_excel(writer, sheet_name="Base", index=False)
        _highlight_total_mismatches(writer, df_base)
    return buffer.getvalue()


def export_to_excel(data, output_dir: str, file_name: str = "") -> str:
    """Exporta exactamente la tabla visible de revision a una hoja Base."""
    output_path = Path(output_dir).expanduser()
    output_path.mkdir(parents=True, exist_ok=True)

    if not file_name.strip():
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"CashCalls_Export_{timestamp}.xlsx"
    elif not file_name.lower().endswith(".xlsx"):
        file_name = f"{file_name}.xlsx"

    file_path = output_path / Path(file_name).name
    df_base = data.copy() if isinstance(data, pd.DataFrame) else pd.DataFrame(data)

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df_base.to_excel(writer, sheet_name="Base", index=False)
        _highlight_total_mismatches(writer, df_base)

    return str(file_path.resolve())


def _highlight_total_mismatches(writer, df: pd.DataFrame) -> None:
    total_field = "Total Current Period"
    if total_field not in df.columns:
        return

    concept_fields = [field for field in CASH_CALL_CONCEPTS if field in df.columns]
    if not concept_fields:
        return

    worksheet = writer.sheets["Base"]
    total_column = list(df.columns).index(total_field) + 1
    for row_index, (_, row) in enumerate(df.iterrows(), start=2):
        total = _float_or_none(row.get(total_field))
        concept_amounts = [_float_or_none(row.get(field)) for field in concept_fields]
        concept_amounts = [amount for amount in concept_amounts if amount is not None]
        if total is None or not concept_amounts:
            continue
        if abs(round(sum(concept_amounts), 2) - total) > 0.05:
            worksheet.cell(row=row_index, column=total_column).fill = TOTAL_MISMATCH_FILL


def _float_or_none(value) -> float | None:
    if value is None or pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None

def export_multiple_sheets_bytes(sheets_data: dict[str, pd.DataFrame]) -> bytes:
    """Genera en memoria un Excel con múltiples hojas."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, df in sheets_data.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    return buffer.getvalue()
