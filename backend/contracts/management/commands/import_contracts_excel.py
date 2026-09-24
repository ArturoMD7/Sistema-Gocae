from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

from contracts.models import (
    Contract,
    Contractor,
    GeologicalData,
    Investment,
    LocationPoint,
    NationalContentRule,
    PlanActivity,
    PlanDocument,
    PmtInfo,
    PmtMethod,
    PmtPenalty,
    PmtStudy,
    PmtSummary,
    PmtWell,
    Resolution,
    StateRevenue,
    TimelineEvent,
)


EMPTY_STRINGS = {"", "n/a", "(en blanco)", "nan", "none", "null"}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return None if value.lower() in EMPTY_STRINGS else value
    return value


def as_text(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_decimal(value):
    value = clean(value)
    if value is None:
        return None
    try:
        return Decimal(str(value).replace(",", ""))
    except (InvalidOperation, ValueError):
        return None


def as_int(value):
    value = as_decimal(value)
    if value is None:
        return None
    return int(value)


def as_date(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


SHEETS = {
    "TABLA CONTRATO": {
        "model": Contract,
        "pk": "contract_id",
        "fields": {
            "ID_Contrato": ("contract_id", as_text),
            "Modalidad": ("modalidad", as_text),
            "Ronda": ("ronda", as_text),
            "Licitacion": ("licitacion", as_text),
            "Area": ("area", as_decimal),
            "Pozos_Comprometidos": ("pozos_comprometidos", as_decimal),
            "Pozos_Terminados": ("pozos_terminados", as_decimal),
            "Operador_Principal": ("operador_principal", as_text),
            "Tipo_Licitante": ("tipo_licitante", as_text),
            "Tipo_Yacimiento": ("tipo_yacimiento", as_text),
            "Estado": ("estado", as_text),
            "Superficie_km2": ("superficie_km2", as_decimal),
            "Fecha_Firma": ("fecha_firma", as_date),
            "Duracion_Años": ("duracion_anios", as_decimal),
            "Vence": ("vence", as_int),
            "Estatus": ("estatus", as_text),
            "Regla_Exploracion": ("regla_exploracion", as_text),
            "No_Registro_Fiducario": ("no_registro_fiducario", as_text),
            "Fecha_Registro_Fiducario": ("fecha_registro_fiducario", as_date),
            "Participacion_Estado": ("participacion_estado", as_decimal),
            "Valor_Regalia_Adicional": ("valor_regalia_adicional", as_decimal),
            "Incremento_Programa_Min": ("incremento_programa_min", as_text),
            "Valor_Ponderado_Prop_Economica": ("valor_ponderado_prop_economica", as_decimal),
            "Resultados_": ("resultados", as_text),
        },
    },
    "LINEA DEL TIEMPO": {
        "model": TimelineEvent,
        "pk": "timeline_id",
        "fields": {
            "ID_Timeline": ("timeline_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Fecha": ("fecha", as_date),
            "Evento": ("evento", as_text),
            "Detalles_Evento": ("detalles_evento", as_text),
        },
    },
    "UBICACION": {
        "model": LocationPoint,
        "pk": "point_id",
        "fields": {
            "ID_Punto": ("point_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Numero_Vertice": ("numero_vertice", as_decimal),
            "Longitud_Decimales": ("longitud_decimales", as_decimal),
            "Latitud_Decimales": ("latitud_decimales", as_decimal),
        },
    },
    "CONTRATISTA": {
        "model": Contractor,
        "pk": "partner_id",
        "fields": {
            "ID_Socio": ("partner_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Empresa_Nombre": ("empresa_nombre", as_text),
            "Tipo_Participacion": ("tipo_participacion", as_text),
            "Porcentaje_Participacion": ("porcentaje_participacion", as_decimal),
            "Pais_Origen": ("pais_origen", as_text),
            "Actividad_Empresarial": ("actividad_empresarial", as_text),
            "Socios_Accionistas_Principales": ("socios_accionistas_principales", as_text),
            "Cotizacion_Bolsa": ("cotizacion_bolsa", as_text),
        },
    },
    "SITUACION ACTUAL": {
        "model": Resolution,
        "pk": "resolution_id",
        "fields": {
            "ID_Resolucion": ("resolution_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Actividad_Regulada": ("actividad_regulada", as_text),
            "Fecha_Presentacion": ("fecha_presentacion", as_date),
            "Fecha_Resolucion": ("fecha_resolucion", as_date),
            "Sentido_Resolucion": ("sentido_resolucion", as_text),
            "Numero_Resolucion": ("numero_resolucion", as_text),
        },
    },
    "PTM RESUMEN": {
        "model": PmtSummary,
        "pk": "pmt_id",
        "fields": {
            "ID_PMT": ("pmt_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Programa_Minimo_Trabajo_UT": ("programa_minimo_trabajo_ut", as_decimal),
            "Incremento_PMT_UT": ("incremento_pmt_ut", as_decimal),
            "Periodo_Adicional_UT": ("periodo_adicional_ut", as_decimal),
            "PMT_Total_Requerido": ("pmt_total_requerido", as_decimal),
            "Acreditadas_Reales": ("acreditadas_reales", as_decimal),
            "Fecha_Limite": ("fecha_limite", as_date),
        },
    },
    "PTM POZOS": {
        "model": PmtWell,
        "pk": "pmt_well_id",
        "fields": {
            "ID_PMT_Pozo": ("pmt_well_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Profundidad_Metros": ("profundidad_metros", as_decimal),
            "Unidades_Trabajo": ("unidades_trabajo", as_decimal),
            "Profundidad de perforación (mbnm)": ("profundidad_perforacion_mbnm", as_text),
            "Tirante de agua (metros) 500-1000": ("tirante_agua_500_1000", as_text),
            "Tirante de agua (metros) >1000-2000": ("tirante_agua_1000_2000", as_text),
            "Tirante de agua (metros) >2000": ("tirante_agua_mayor_2000", as_text),
        },
    },
    "PTM INFO": {
        "model": PmtInfo,
        "pk": "pmt_info_id",
        "fields": {
            "ID_PMT_Info": ("pmt_info_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Actividad_Exploratoria": ("actividad_exploratoria", as_text),
            "Descripcion_Actividad": ("descripcion_actividad", as_text),
            "Unidad_Medida": ("unidad_medida", as_text),
            "Unidades_Trabajo": ("unidades_trabajo", as_decimal),
        },
    },
    "PTM ESTUDIOS": {
        "model": PmtStudy,
        "pk": "pmt_study_id",
        "fields": {
            "ID_PMT_Estudio": ("pmt_study_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Actividad_Exploratoria": ("actividad_exploratoria", as_text),
            "Descripcion_Actividad": ("descripcion_actividad", as_text),
            "Unidad_Medida": ("unidad_medida", as_text),
            "Unidades_Trabajo": ("unidades_trabajo", as_decimal),
            "Estudios exploratorios": ("estudios_exploratorios", as_text),
        },
    },
    "PTM METODOS": {
        "model": PmtMethod,
        "pk": "pmt_method_id",
        "fields": {
            "ID_PMT_Metodo": ("pmt_method_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Actividad_Exploratoria": ("actividad_exploratoria", as_text),
            "Descripcion_Actividad": ("descripcion_actividad", as_text),
            "Unidad_Medida": ("unidad_medida", as_text),
            "Unidades_Trabajo": ("unidades_trabajo", as_decimal),
        },
    },
    "PTM PENALIZACION": {
        "model": PmtPenalty,
        "pk": "penalty_id",
        "fields": {
            "PMT_PENALIZACIONES": ("penalty_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Rango_Brent_Minimo_USD": ("rango_brent_minimo_usd", as_decimal),
            "Rango_Brent_Maximo_USD": ("rango_brent_maximo_usd", as_decimal),
            "Valor_Por_Unidad_USD": ("valor_por_unidad_usd", as_decimal),
        },
    },
    "REGLAS CONTENIDO NACIONAL": {
        "model": NationalContentRule,
        "pk": "rule_id",
        "fields": {
            "ID_Regla_CN": ("rule_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Etapa": ("etapa", as_text),
            "Porcentaje_Base": ("porcentaje_base", as_decimal),
            "Descripcion_Regla": ("descripcion_regla", as_text),
        },
    },
    "PLANES": {
        "model": PlanDocument,
        "pk": "document_id",
        "fields": {
            "ID_Documento": ("document_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Tipo_Documento": ("tipo_documento", as_text),
            "Fecha_Documento": ("fecha_documento", as_text),
        },
    },
    "PLANES ACTIVIDADES": {
        "model": PlanActivity,
        "pk": "activity_id",
        "fields": {
            "ID_Actividad": ("activity_id", as_text),
            "ID_Documento": ("document_id", as_text),
            "Sub_Actividad": ("sub_actividad", as_text),
            "Tarea / Descripción": ("tarea_descripcion", as_text),
            "Año_Inicio": ("anio_inicio", as_int),
            "Año_Fin": ("anio_fin", as_int),
            "Mes_Inicio": ("mes_inicio", as_text),
            "Mes_Fin": ("mes_fin", as_text),
            "Cantidad": ("cantidad", as_decimal),
            "UT_Aporta": ("ut_aporta", as_decimal),
        },
    },
    "INVERSIONES EJERCIDAS": {
        "model": Investment,
        "pk": "investment_id",
        "fields": {
            "ID_Inversion": ("investment_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Año": ("anio", as_int),
            "Mes": ("mes", as_text),
            "Fase_Actividad": ("fase_actividad", as_text),
            "Monto_Aprobado_USD": ("monto_aprobado_usd", as_decimal),
            "Monto_Ejercido_USD": ("monto_ejercido_usd", as_decimal),
        },
    },
    "INGRESOS ESTADO": {
        "model": StateRevenue,
        "pk": "revenue_id",
        "fields": {
            "ID_Ingreso": ("revenue_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Año": ("anio", as_int),
            "Mes": ("mes", as_text),
            "Concepto_Pago": ("concepto_pago", as_text),
            "Monto_Pagado_MXN": ("monto_pagado_mxn", as_decimal),
            "Monto_Pagado_USD": ("monto_pagado_usd", as_decimal),
        },
    },
    "DATOS GEOLOGICOS": {
        "model": GeologicalData,
        "pk": "geological_id",
        "fields": {
            "ID_DatGeo": ("geological_id", as_text),
            "ID_Contrato": ("contract_id", as_text),
            "Provincia_Petrolera": ("provincia_petrolera", as_text),
            "Provincia_Geologica": ("provincia_geologica", as_text),
            "Superficie_AContractual_km2": ("superficie_acontractual_km2", as_decimal),
            "Cobertura_sismica_3D": ("cobertura_sismica_3d", as_decimal),
            "Edad_Play": ("edad_play", as_text),
            "Litologias": ("litologias", as_text),
            "Hidrocarburo_Esperado": ("hidrocarburo_esperado", as_text),
        },
    },
}


DELETE_ORDER = [
    GeologicalData,
    StateRevenue,
    Investment,
    PlanActivity,
    PlanDocument,
    NationalContentRule,
    PmtPenalty,
    PmtMethod,
    PmtStudy,
    PmtInfo,
    PmtWell,
    PmtSummary,
    Resolution,
    Contractor,
    LocationPoint,
    TimelineEvent,
    Contract,
]


class Command(BaseCommand):
    help = "Importa la propuesta de BD grupal de contratos GOCAE desde Excel."

    def add_arguments(self, parser):
        parser.add_argument("excel_path")
        parser.add_argument("--clear", action="store_true", help="Borra datos de contratos antes de importar.")

    @transaction.atomic
    def handle(self, *args, **options):
        if load_workbook is None:
            raise CommandError("Falta openpyxl. Instala requirements.txt antes de importar.")

        workbook = load_workbook(options["excel_path"], data_only=True, read_only=True)

        if options["clear"]:
            for model in DELETE_ORDER:
                model.objects.all().delete()

        totals = {}
        for sheet_name, config in SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                self.stdout.write(self.style.WARNING(f"Hoja no encontrada: {sheet_name}"))
                continue

            ws = workbook[sheet_name]
            headers = [as_text(value) for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
            index = {header: pos for pos, header in enumerate(headers) if header}

            created = updated = skipped = 0
            model = config["model"]
            pk_field = config["pk"]
            field_map = config["fields"]

            for row in ws.iter_rows(min_row=2, values_only=True):
                payload = {}
                for excel_header, (field_name, converter) in field_map.items():
                    if excel_header not in index:
                        continue
                    payload[field_name] = converter(row[index[excel_header]])

                pk_value = payload.get(pk_field)
                if not pk_value:
                    skipped += 1
                    continue

                defaults = {key: value for key, value in payload.items() if key != pk_field}
                _, was_created = model.objects.update_or_create(**{pk_field: pk_value}, defaults=defaults)
                if was_created:
                    created += 1
                else:
                    updated += 1

            totals[sheet_name] = {"created": created, "updated": updated, "skipped": skipped}

        for sheet, counts in totals.items():
            self.stdout.write(
                f"{sheet}: {counts['created']} creados, {counts['updated']} actualizados, {counts['skipped']} omitidos"
            )
